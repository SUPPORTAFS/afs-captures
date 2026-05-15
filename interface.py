import streamlit as st
import cv2
import os
import numpy as np
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


# =========================
# OUTILS
# =========================

def clean_folder(folder):
    os.makedirs(folder, exist_ok=True)
    for file in os.listdir(folder):
        try:
            os.remove(os.path.join(folder, file))
        except:
            pass


def analyse_video(video_path, output_folder, prefix, max_captures=6):
    clean_folder(output_folder)

    interval_seconds = 3
    blur_threshold = 200
    difference_threshold = 15

    video = cv2.VideoCapture(video_path)
    fps = video.get(cv2.CAP_PROP_FPS)

    if fps == 0:
        fps = 25

    frame_interval = int(fps * interval_seconds)
    frame_count = 0
    capture_count = 0
    last_saved_gray = None

    while True:
        success, frame = video.read()

        if not success:
            break

        if capture_count >= max_captures:
            break

        if frame_count % frame_interval == 0:
            gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)

            sharpness = cv2.Laplacian(gray, cv2.CV_64F).var()

            if sharpness < blur_threshold:
                frame_count += 1
                continue

            if last_saved_gray is not None:
                difference = cv2.absdiff(gray, last_saved_gray)
                mean_difference = np.mean(difference)

                if mean_difference < difference_threshold:
                    frame_count += 1
                    continue

            image_path = os.path.join(
                output_folder,
                f"{prefix}_{capture_count + 1}.jpg"
            )

            cv2.imwrite(image_path, frame)

            last_saved_gray = gray
            capture_count += 1

        frame_count += 1

    video.release()
    return capture_count


def reformuler_client(texte, rubrique):
    if not texte.strip():
        return ""

    texte = texte.strip()

    if rubrique == "avant":
        return (
            "Lors de l’inspection avant intervention, nous constatons les éléments suivants : "
            + texte
            + "."
        )

    if rubrique == "travaux":
        return (
            "L’intervention réalisée a consisté à effectuer les opérations suivantes : "
            + texte
            + "."
        )

    if rubrique == "apres":
        return (
            "Après intervention, le contrôle met en évidence les éléments suivants : "
            + texte
            + "."
        )

    if rubrique == "conclusion":
        return (
            "En conclusion, "
            + texte
            + "."
        )

    return texte


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(18)


def add_heading(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)


def add_text(doc, text):
    if not text:
        text = "Non renseigné."

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    run.font.size = Pt(10)


def add_image(doc, image_path, width=5.8):
    if os.path.exists(image_path):
        doc.add_picture(image_path, width=Inches(width))


def generate_word(data, selected_before, selected_after):
    doc_path = "rapport_intervention_afs.docx"

    doc = Document()

    # PAGE DE GARDE
    if os.path.exists("Logo.png"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture("Logo.png", width=Inches(2.5))

    add_title(doc, "RAPPORT D’INTERVENTION")
    add_title(doc, "INSPECTION CAMÉRA / CONTRÔLE RÉSEAU")

    doc.add_paragraph("")
    add_text(doc, f"Site : {data['site']}")
    add_text(doc, f"Client : {data['client']}")
    add_text(doc, f"Adresse : {data['adresse']}")
    add_text(doc, f"Intervenant : {data['intervenant']}")
    add_text(doc, f"Date : {data['date']}")
    add_text(doc, f"Référence : {data['reference']}")
    add_text(doc, f"Objet : {data['objet']}")

    doc.add_page_break()

    add_heading(doc, "1. Constat et objet de l’intervention")
    add_text(doc, reformuler_client(data["constat"], "avant"))

    doc.add_paragraph("")

    add_heading(doc, "2. Matériels utilisés")
    add_text(
        doc,
        "Afin de réaliser un contrôle des canalisations, nous utilisons un système "
        "d’inspection caméra permettant de visualiser l’état intérieur des réseaux. "
        "Cet équipement permet d’identifier les anomalies telles que dépôts, fissures, "
        "déformations, obturations, contre-pentes ou défauts d’étanchéité."
    )

    doc.add_paragraph("")

    add_heading(doc, "3. Méthodologie d’intervention")
    add_text(
        doc,
        "Après identification du réseau concerné, l’inspection est réalisée par introduction "
        "de la caméra dans la canalisation. Les observations sont relevées au fur et à mesure "
        "de la progression afin de documenter l’état du réseau avant et après intervention."
    )

    doc.add_paragraph("")

    add_heading(doc, "4. État avant intervention")
    add_text(doc, reformuler_client(data["etat_avant"], "avant"))

    for image_path in selected_before:
        add_image(doc, image_path)
        doc.add_paragraph("Capture avant intervention")

    doc.add_paragraph("")

    add_heading(doc, "5. Intervention réalisée")
    add_text(doc, reformuler_client(data["travaux"], "travaux"))

    doc.add_paragraph("")

    add_heading(doc, "6. État après intervention")
    add_text(doc, reformuler_client(data["etat_apres"], "apres"))

    for image_path in selected_after:
        add_image(doc, image_path)
        doc.add_paragraph("Capture après intervention")

    doc.add_paragraph("")

    add_heading(doc, "7. Conclusion")
    add_text(doc, reformuler_client(data["conclusion"], "conclusion"))

    doc.save(doc_path)
    return doc_path


# =========================
# CONFIG
# =========================

st.set_page_config(
    page_title="AFS PLOMBERIE CVC",
    layout="wide"
)

st.image("Logo.png", width=220)
st.title("Générateur de rapport d’intervention")

if st.button("Réinitialiser les captures"):
    clean_folder("captures_avant")
    clean_folder("captures_apres")

    if os.path.exists("video_avant.mp4"):
        os.remove("video_avant.mp4")

    if os.path.exists("video_apres.mp4"):
        os.remove("video_apres.mp4")

    st.success("Captures réinitialisées. Vous pouvez charger de nouvelles vidéos.")
    st.rerun()


# =========================
# FORMULAIRE RAPPORT
# =========================

st.header("Informations du rapport")

col1, col2 = st.columns(2)

with col1:
    site = st.text_input("Site")
    client = st.text_input("Client")
    adresse = st.text_input("Adresse")
    intervenant = st.text_input("Intervenant")

with col2:
    date = st.text_input("Date")
    reference = st.text_input("Nos références")
    objet = st.text_input("Objet de l’intervention")

constat = st.text_area("Constat / mission", height=100)

etat_avant = st.text_area(
    "État AVANT intervention - notes technicien",
    height=120,
    placeholder="Ex : canalisation encrassée, dépôts importants, écoulement difficile..."
)

travaux = st.text_area(
    "Ce qui a été fait - notes technicien",
    height=120,
    placeholder="Ex : passage caméra, curage haute pression, contrôle final..."
)

etat_apres = st.text_area(
    "État APRÈS intervention - notes technicien",
    height=120,
    placeholder="Ex : écoulement rétabli, canalisation propre, plus d’obstruction visible..."
)

conclusion = st.text_area(
    "Conclusion / préconisations",
    height=120,
    placeholder="Ex : intervention concluante, réseau fonctionnel, surveillance recommandée..."
)


# =========================
# UPLOAD VIDEOS
# =========================

st.header("Vidéos")

uploaded_video_before = st.file_uploader(
    "Vidéo AVANT intervention",
    type=["mp4"]
)

uploaded_video_after = st.file_uploader(
    "Vidéo APRÈS intervention",
    type=["mp4"]
)

if uploaded_video_before is not None and uploaded_video_after is not None:
    video_before_path = "video_avant.mp4"

    with open(video_before_path, "wb") as f:
        f.write(uploaded_video_before.read())

    video_after_path = "video_apres.mp4"

    with open(video_after_path, "wb") as f:
        f.write(uploaded_video_after.read())

    st.success("Vidéos chargées")

    if st.button("Sélectionner les captures à intégrer dans le rapport"):
        with st.spinner("Analyse en cours..."):
            count_before = analyse_video(
                video_before_path,
                "captures_avant",
                "avant",
                max_captures=6
            )

            count_after = analyse_video(
                video_after_path,
                "captures_apres",
                "apres",
                max_captures=6
            )

        st.success(
            f"Analyse terminée : {count_before} captures AVANT et {count_after} captures APRÈS"
        )


# =========================
# SELECTION CAPTURES AVANT
# =========================

selected_before = []

if os.path.exists("captures_avant"):
    st.header("Captures AVANT à intégrer au rapport")

    files_before = sorted(os.listdir("captures_avant"))
    cols_before = st.columns(3)

    for index, file in enumerate(files_before):
        if not file.endswith(".jpg"):
            continue

        image_path = os.path.join("captures_avant", file)

        with cols_before[index % 3]:
            st.image(image_path, use_container_width=True)

            selected = st.checkbox(
                "Intégrer au rapport",
                key=f"before_{file}"
            )

            if selected:
                selected_before.append(image_path)


# =========================
# SELECTION CAPTURES APRES
# =========================

selected_after = []

if os.path.exists("captures_apres"):
    st.header("Captures APRÈS à intégrer au rapport")

    files_after = sorted(os.listdir("captures_apres"))
    cols_after = st.columns(3)

    for index, file in enumerate(files_after):
        if not file.endswith(".jpg"):
            continue

        image_path = os.path.join("captures_apres", file)

        with cols_after[index % 3]:
            st.image(image_path, use_container_width=True)

            selected = st.checkbox(
                "Intégrer au rapport",
                key=f"after_{file}"
            )

            if selected:
                selected_after.append(image_path)


# =========================
# GENERATION WORD
# =========================

st.header("Génération du rapport Word")

if st.button("Générer le rapport Word"):
    data = {
        "site": site,
        "client": client,
        "adresse": adresse,
        "intervenant": intervenant,
        "date": date,
        "reference": reference,
        "objet": objet,
        "constat": constat,
        "etat_avant": etat_avant,
        "travaux": travaux,
        "etat_apres": etat_apres,
        "conclusion": conclusion,
    }

    docx_path = generate_word(
        data,
        selected_before,
        selected_after
    )

    st.success("Rapport Word généré")

    with open(docx_path, "rb") as file:
        st.download_button(
            label="Télécharger le rapport Word",
            data=file,
            file_name="rapport_intervention_afs.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )